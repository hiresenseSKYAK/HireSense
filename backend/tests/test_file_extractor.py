import io
import unittest

from docx import Document

from services.file_extractor import extract_resume_text


def minimal_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode("ascii"))
    return bytes(output)


class FileExtractorTests(unittest.TestCase):
    def test_extracts_text_from_pdf_bytes(self):
        text = extract_resume_text("resume.pdf", minimal_pdf("Ada Lovelace - Python Engineer"))
        self.assertIn("Ada Lovelace", text)
        self.assertIn("Python Engineer", text)

    def test_extracts_text_from_docx_bytes(self):
        document = Document()
        document.add_paragraph("Grace Hopper")
        document.add_paragraph("Skills: Python, SQL")
        buffer = io.BytesIO()
        document.save(buffer)

        text = extract_resume_text("resume.docx", buffer.getvalue())
        self.assertIn("Grace Hopper", text)
        self.assertIn("Python, SQL", text)

    def test_rejects_unsupported_files(self):
        with self.assertRaisesRegex(ValueError, "PDF and DOCX"):
            extract_resume_text("resume.txt", b"plain text")


if __name__ == "__main__":
    unittest.main()
